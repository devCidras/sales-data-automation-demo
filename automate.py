"""Excel/CSV sales workflow automation demo.

Problem it solves: a small business exports raw sales data (Excel or CSV)
that is inconsistently formatted, contains duplicates, missing values and
mixed date/currency formats. Someone then spends time by hand cleaning it up
in Excel and building a summary for the owner.

This script does that automatically:
1. Loads the raw export (.csv or .xlsx).
2. Validates and cleans it, separating unusable rows instead of silently
   dropping them.
3. Computes business KPIs (revenue, top products, monthly trend, returns...).
4. Writes a formatted Excel workbook (clean data + KPI summary + rejected
   rows) and a one-page PDF report with charts.

Usage:
    python automate.py --input data/raw_sales.csv --output-dir output
"""

from __future__ import annotations

import argparse
import re
import time
from dataclasses import dataclass
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CURRENCY_RE = re.compile(r"[^\d,.\-]")
DATE_FORMATS = ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y")


@dataclass
class Kpis:
    total_revenue: float
    total_returns: float
    net_revenue: float
    total_units_sold: int
    num_line_items: int
    avg_line_item_value: float
    revenue_by_category: pd.DataFrame
    revenue_by_region: pd.DataFrame
    top_products: pd.DataFrame
    monthly_revenue: pd.DataFrame


def load_raw(path: Path) -> pd.DataFrame:
    if path.suffix.lower() == ".csv":
        return pd.read_csv(path, dtype=str)
    return pd.read_excel(path, dtype=str)


def _parse_date(value: str):
    if not isinstance(value, str) or not value.strip():
        return pd.NaT
    for fmt in DATE_FORMATS:
        try:
            return pd.to_datetime(value.strip(), format=fmt)
        except ValueError:
            continue
    return pd.to_datetime(value, errors="coerce", dayfirst=True)


def _parse_price(value: str) -> float:
    """Parses a currency string, handling both US (1,234.56) and EU
    (1.234,56) thousands/decimal separator conventions."""
    if not isinstance(value, str) or not value.strip():
        return float("nan")
    cleaned = CURRENCY_RE.sub("", value).strip()
    if not cleaned:
        return float("nan")

    has_comma = "," in cleaned
    has_dot = "." in cleaned
    if has_comma and has_dot:
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")
    elif has_comma:
        decimals = cleaned[cleaned.rfind(",") + 1 :]
        if len(decimals) == 3 and decimals.isdigit():
            cleaned = cleaned.replace(",", "")
        else:
            cleaned = cleaned.replace(",", ".")

    try:
        return float(cleaned)
    except ValueError:
        return float("nan")


def clean_and_validate(raw: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    df = raw.copy()

    df["Product"] = df["Product"].fillna("").astype(str).str.strip().str.title()
    df["Category"] = df["Category"].fillna("").astype(str).str.strip().str.title()
    df["Region"] = df["Region"].fillna("").astype(str).str.strip().str.title()
    df["Region"] = df["Region"].replace({"": "Unknown"})
    df["Customer"] = df["Customer"].fillna("").astype(str).str.strip()
    df["Customer"] = df["Customer"].replace({"": "Unknown"})

    df["Date"] = df["Date"].apply(_parse_date)
    df["Unit Price"] = df["Unit Price"].apply(_parse_price)
    df["Quantity"] = pd.to_numeric(df["Quantity"], errors="coerce")

    before_dupes = len(df)
    df = df.drop_duplicates()
    duplicates_removed = before_dupes - len(df)

    missing_date = df["Date"].isna()
    missing_price = df["Unit Price"].isna()
    missing_qty = df["Quantity"].isna() | (df["Quantity"] == 0)

    reasons = []
    for md, mp, mq in zip(missing_date, missing_price, missing_qty):
        row_reasons = []
        if md:
            row_reasons.append("invalid date")
        if mp:
            row_reasons.append("invalid price")
        if mq:
            row_reasons.append("invalid quantity")
        reasons.append(", ".join(row_reasons))

    df["_rejection_reason"] = reasons
    invalid_mask = missing_date | missing_price | missing_qty

    rejected = df[invalid_mask].copy()
    clean = df[~invalid_mask].drop(columns=["_rejection_reason"]).copy()

    clean["Total Value"] = (clean["Quantity"] * clean["Unit Price"]).round(2)
    clean["Type"] = clean["Quantity"].apply(lambda q: "Return" if q < 0 else "Sale")
    clean = clean.sort_values("Date").reset_index(drop=True)

    stats = {
        "original_rows": len(raw),
        "duplicates_removed": duplicates_removed,
        "rejected_rows": len(rejected),
        "valid_rows": len(clean),
    }
    return clean, rejected, stats


def compute_kpis(clean: pd.DataFrame) -> Kpis:
    sales = clean[clean["Type"] == "Sale"]
    returns = clean[clean["Type"] == "Return"]

    total_revenue = round(sales["Total Value"].sum(), 2)
    total_returns = round(-returns["Total Value"].sum(), 2)
    net_revenue = round(total_revenue - total_returns, 2)
    total_units_sold = int(sales["Quantity"].sum())
    num_line_items = len(sales)
    avg_line_item_value = round(total_revenue / num_line_items, 2) if num_line_items else 0.0

    revenue_by_category = (
        sales.groupby("Category")["Total Value"].sum().round(2).sort_values(ascending=False).reset_index()
    )
    revenue_by_region = (
        sales.groupby("Region")["Total Value"].sum().round(2).sort_values(ascending=False).reset_index()
    )
    top_products = (
        sales.groupby("Product")["Total Value"]
        .sum()
        .round(2)
        .sort_values(ascending=False)
        .head(5)
        .reset_index()
    )

    monthly = sales.copy()
    monthly["Month"] = monthly["Date"].dt.to_period("M").astype(str)
    monthly_revenue = monthly.groupby("Month")["Total Value"].sum().round(2).reset_index()

    return Kpis(
        total_revenue=total_revenue,
        total_returns=total_returns,
        net_revenue=net_revenue,
        total_units_sold=total_units_sold,
        num_line_items=num_line_items,
        avg_line_item_value=avg_line_item_value,
        revenue_by_category=revenue_by_category,
        revenue_by_region=revenue_by_region,
        top_products=top_products,
        monthly_revenue=monthly_revenue,
    )


HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autofit(ws, df: pd.DataFrame) -> None:
    for i, col in enumerate(df.columns, start=1):
        text_lengths = df[col].fillna("").astype(str).map(len)
        width = max(len(str(col)), text_lengths.max() if len(df) else 0) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 40)


def export_excel(clean: pd.DataFrame, rejected: pd.DataFrame, kpis: Kpis, stats: dict, out_path: Path) -> None:
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        clean.to_excel(writer, sheet_name="Clean Data", index=False)
        rejected.to_excel(writer, sheet_name="Rejected Rows", index=False)

        summary_rows = [
            ("Rows in original file", stats["original_rows"]),
            ("Duplicates removed", stats["duplicates_removed"]),
            ("Rejected rows (invalid data)", stats["rejected_rows"]),
            ("Valid rows processed", stats["valid_rows"]),
            ("", ""),
            ("Total revenue (sales)", kpis.total_revenue),
            ("Value in returns", kpis.total_returns),
            ("Net revenue", kpis.net_revenue),
            ("Units sold", kpis.total_units_sold),
            ("Number of line items", kpis.num_line_items),
            ("Average line item value", kpis.avg_line_item_value),
        ]
        summary_df = pd.DataFrame(summary_rows, columns=["Metric", "Value"])
        summary_df.to_excel(writer, sheet_name="KPI Summary", index=False, startrow=0)

        row = len(summary_df) + 3
        kpis.revenue_by_category.to_excel(writer, sheet_name="KPI Summary", index=False, startrow=row)

        row += len(kpis.revenue_by_category) + 3
        kpis.revenue_by_region.to_excel(writer, sheet_name="KPI Summary", index=False, startrow=row)

        row += len(kpis.revenue_by_region) + 3
        kpis.top_products.to_excel(writer, sheet_name="KPI Summary", index=False, startrow=row)

        for sheet_name, df in (
            ("Clean Data", clean),
            ("Rejected Rows", rejected),
            ("KPI Summary", summary_df),
        ):
            ws = writer.sheets[sheet_name]
            ws.freeze_panes = "A2"
            _style_header(ws, len(df.columns))
            _autofit(ws, df)
            ws.auto_filter.ref = ws.dimensions


def _save_charts(kpis: Kpis, charts_dir: Path) -> dict[str, Path]:
    charts_dir.mkdir(parents=True, exist_ok=True)
    paths = {}

    fig, ax = plt.subplots(figsize=(5, 2.8))
    ax.plot(kpis.monthly_revenue["Month"], kpis.monthly_revenue["Total Value"], marker="o", color="#1F4E78")
    ax.set_title("Monthly revenue")
    ax.set_ylabel("EUR")
    fig.tight_layout()
    monthly_path = charts_dir / "monthly_revenue.png"
    fig.savefig(monthly_path, dpi=150)
    plt.close(fig)
    paths["monthly"] = monthly_path

    fig, ax = plt.subplots(figsize=(5, 2.8))
    ax.barh(kpis.top_products["Product"], kpis.top_products["Total Value"], color="#2E75B6")
    ax.set_title("Top 5 products by revenue")
    ax.invert_yaxis()
    fig.tight_layout()
    top_path = charts_dir / "top_products.png"
    fig.savefig(top_path, dpi=150)
    plt.close(fig)
    paths["top_products"] = top_path

    return paths


def export_pdf(kpis: Kpis, stats: dict, charts: dict[str, Path], out_path: Path) -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 12, "Sales Report", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(100, 100, 100)
    pdf.cell(
        0,
        6,
        f"Automatically generated - {pd.Timestamp.now():%Y-%m-%d %H:%M}",
        new_x=XPos.LMARGIN,
        new_y=YPos.NEXT,
    )
    pdf.set_text_color(0, 0, 0)
    pdf.ln(4)

    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Key metrics", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 11)
    kpi_lines = [
        f"Total revenue: EUR {kpis.total_revenue:,.2f}",
        f"Net revenue (after returns): EUR {kpis.net_revenue:,.2f}",
        f"Units sold: {kpis.total_units_sold}",
        f"Number of line items: {kpis.num_line_items}",
        f"Average line item value: EUR {kpis.avg_line_item_value:,.2f}",
    ]
    for line in kpi_lines:
        pdf.cell(0, 7, line, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(3)

    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Data quality", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "", 11)
    quality_lines = [
        f"Rows in original file: {stats['original_rows']}",
        f"Duplicates removed automatically: {stats['duplicates_removed']}",
        f"Invalid rows flagged for review: {stats['rejected_rows']}",
        f"Valid rows used in the report: {stats['valid_rows']}",
    ]
    for line in quality_lines:
        pdf.cell(0, 7, line, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(4)

    chart_y = pdf.get_y() + 2
    chart_w = 88
    pdf.image(str(charts["monthly"]), x=10, y=chart_y, w=chart_w)
    pdf.image(str(charts["top_products"]), x=112, y=chart_y, w=chart_w)

    pdf.output(str(out_path))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=Path("data/raw_sales.csv"))
    parser.add_argument("--output-dir", type=Path, default=Path("output"))
    args = parser.parse_args()

    args.output_dir.mkdir(parents=True, exist_ok=True)
    start = time.perf_counter()

    raw = load_raw(args.input)
    clean, rejected, stats = clean_and_validate(raw)
    kpis = compute_kpis(clean)

    excel_path = args.output_dir / "processed_sales.xlsx"
    pdf_path = args.output_dir / "sales_report.pdf"
    charts_dir = args.output_dir / "_charts"

    export_excel(clean, rejected, kpis, stats, excel_path)
    charts = _save_charts(kpis, charts_dir)
    export_pdf(kpis, stats, charts, pdf_path)

    elapsed = time.perf_counter() - start
    manual_estimate_minutes = max(30, stats["original_rows"] // 8)

    print("Processing complete.")
    print(f"  Original rows:        {stats['original_rows']}")
    print(f"  Duplicates removed:   {stats['duplicates_removed']}")
    print(f"  Rejected rows:        {stats['rejected_rows']}")
    print(f"  Valid rows:           {stats['valid_rows']}")
    print(f"  Net revenue:          EUR {kpis.net_revenue:,.2f}")
    print(f"  Line items:           {kpis.num_line_items}")
    print(f"  Excel generated:      {excel_path}")
    print(f"  PDF generated:        {pdf_path}")
    print(f"  Execution time:       {elapsed:.2f} s")
    print(f"  Estimated manual time: ~{manual_estimate_minutes} min")


if __name__ == "__main__":
    main()
